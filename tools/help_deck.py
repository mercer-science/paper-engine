"""Build help/paper_engine_overview.pptx - the flow chart of how the aids interact.

    python tools/help_deck.py build                 # -> help/paper_engine_overview.pptx
    python tools/help_deck.py build --no-animation  # same deck, no timing block
    python tools/help_deck.py check                 # re-open it and measure what is in it
    python tools/help_deck.py graph --json          # the stage graph, as data

WHY THIS IS A GENERATOR AND NOT A HAND-DRAWN DECK
The module list changes. A hand-made deck goes stale silently while still
looking authoritative, which is the same failure the toolkit spends most of its
effort on elsewhere - so the stage graph lives in STAGES below, as data, and a
new module is one entry and a rebuild.

Spec: specs/orchestration.md 5.

THE ANIMATION IS INJECTED AS RAW XML, AND THAT IS THE RISKY PART.
python-pptx has no animation API, so `build` writes a <p:timing> block by hand.
Malformed timing XML is one of the few things PowerPoint refuses to open rather
than degrading - hence `--no-animation`, which is the fallback if a future
PowerPoint ever objects, and `check`, which re-parses every part of the written
file before anybody opens it.

Every slide is legible with animations ignored: what the timing block buys is
the order things are revealed in, never the content.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import zipfile

# Windows consoles default to cp1252 and the deck text carries en dashes.
for _stream in (sys.stdout, sys.stderr):
    _reconfigure = getattr(_stream, "reconfigure", None)
    if _reconfigure is not None:
        try:
            _reconfigure(encoding="utf-8", errors="replace")
        except ValueError:
            pass

try:
    from lxml import etree
except ImportError:
    sys.exit("help_deck.py requires 'lxml'. Install with: "
             "python -m pip install lxml")

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("help_deck.py requires 'python-pptx'. Install with: "
             "python -m pip install python-pptx")

TOOLKIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(TOOLKIT, "help", "paper_engine_overview.pptx")

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
# Printed in greyscale as often as not, so the four stage colours are separated
# by lightness as well as by hue - a legend that only works in colour is a
# legend that stops working on the handout.

INK = "1A1A1A"
MUTED = "5B6570"
RULE = "C7CDD4"
PAPER = "FFFFFF"

STAGE_FILL = {
    "A": "E8F0F8", "B": "EAF2EA", "C": "F2EDF7", "D": "FCF0E4",
}
STAGE_LINE = {
    "A": "2E5E8E", "B": "3D6B45", "C": "6A4A8C", "D": "A9601F",
}
ENGINE_FILL = "F1F3F5"
ENGINE_LINE = "8B949E"
LOOP_LINE = "B3341E"
BLIND_FILL = "FFF6E5"

# ---------------------------------------------------------------------------
# The stage graph
# ---------------------------------------------------------------------------
# One entry per stage. `agents` is what the stage calls, in the order it calls
# them; `engines` is the CLI underneath. Read by `graph`, by the slide builders
# below, and by nothing else - help/instructions.md is written by hand
# against the same source of truth (specs/orchestration.md 3.2).

STAGES = [
    {
        "key": "A",
        "title": "Idea Generation",
        "sub": "Before There Is Any Data",
        "agents": [
            ("idea-generation", "six staged rounds: sweep the literature, "
                                "then ask, then 3-5 gaps, then feasibility, "
                                "then the float set"),
            ("setup-project-directory", "scaffolds plan/ data/ drafts/ "
                                        "obsolete/, the R float pipeline, and "
                                        "mock rows shaped like the "
                                        "hypothesis"),
        ],
        "engines": ["scholar.py", "pubmed.py", "structure.py", "idea.py",
                    "scaffold.py"],
        "out": "an Ideas/ folder, a new project, or an additive merge into one "
               "that already has work in it - and figures that exist as claims",
    },
    {
        "key": "B",
        "title": "Research Work",
        # The subtitle used to read "The Aids Get Out Of The Way", which stopped
        # being true the moment a skill stood inside this stage. They still stay
        # out of the bench work; they meet you again at the analysis.
        "sub": "They Stay Out Of The Bench Work",
        "agents": [
            ("analysis", "works out which tests the design supports, fills in "
                         "analysis.R, runs it"),
            ("create-graphic-figure", "a PowerPoint panel for art that is not "
                                      "a chart"),
            ("refine-figure", "edit the float script, render, hold the "
                              "image against the claim"),
        ],
        "engines": ["graphic_figure.py", "scaffold.py float lint",
                    "R pipeline"],
        "out": "real data in data/raw/, images in data/raw_images/, numbers in "
               "data/analysis/analysis.md, floats that render",
    },
    {
        "key": "C",
        "title": "The Writing Engine",
        "sub": "Iterative, And It Learns From You",
        "agents": [
            ("writing-engine", "15 modules, 5 presets at 7/10/13/8/7 agent "
                               "calls; run all of it, one section, or one "
                               "module"),
        ],
        "engines": ["manuscript.py", "prose.py", "learn.py", "docx_edits.py"],
        "out": "manuscript_rN.docx carrying its own outstanding list, plus "
               "a report per check that came back clean",
    },
    {
        "key": "D",
        "title": "Submission",
        "sub": "Everything That Is Not The Manuscript",
        "agents": [
            ("submission-package", "cover letter, title page, statements, "
                                   "reviewers, checklist"),
            ("toc-graphic", "the graphical abstract, sized the journal's way"),
            ("respond-to-reviewers", "one response per point, from the ledger"),
            ("ai-disclosure", "drafts the declaration under one flag - "
                              "YOURS to confirm, not compose; --from-ledger "
                              "adds only what the ledger says it did"),
        ],
        "engines": ["manuscript.py", "toc_graphic.py"],
        "out": "a submission folder that says what is still unanswered, and "
               "refuses to look finished when it is not",
    },
]

# Stage A, as the five things that actually happen, IN THE ORDER THE ENGINE
# RUNS THEM. That order is not the order people list them in: the project
# directory is scaffolded BEFORE there is mock data, because `idea.py mock`
# writes into `data/mock_data/` and needs the tree and the data contract to
# exist first, and the draft figures come last because they are rendered FROM
# the mock data by `scaffold.py mock-floats`.
IDEA_STEPS = [
    ("Research Articles",
     "scholar.py, six indexes - never PubMed alone. 3-8 papers land in "
     "plan/relevant_literature/ with a summary each. Every citation is "
     "verified before it reaches a file; none of them come from memory."),
    ("Idea Discussion",
     "six staged rounds: what you want to write about -> the sweep -> 2-4 "
     "questions asked AFTER the sweep -> 3-5 candidate gaps with their papers "
     "-> feasibility, where “idk” is a real answer -> 2-3 ways to "
     "write it, and the float set."),
    ("Project Directory Set Up",
     "setup-project-directory scaffolds plan/ data/ drafts/ obsolete/ and the "
     "R float pipeline. Six questions, every one with a default. It never "
     "overwrites anything - run it on a folder you already have and it "
     "repairs rather than replaces."),
    ("Mock Data",
     "data/mock_data/*_mock.csv, generated from the data contract the idea "
     "session wrote. The _mock suffix IS the safety mechanism: a mock float "
     "is watermarked and the coauthor .docx is refused while any input is "
     "mock. Adjust the generator and re-run as the real design settles."),
    ("Draft Figures",
     "one float, one folder: plan/figures/FigNN/figure.R. Flip its BUILT "
     "guard to TRUE, point it at the mock CSV, run Rscript "
     "plan/render_all.R - and the paper's central figure is on screen in an "
     "afternoon. Point it at data/raw/ later and the MOCK stamp goes away."),
]

# Stage B, as the files the writing engine actually reads. This list is
# WEIGHTABLE_SOURCES in manuscript.py plus the two submission files, and it is
# the answer to "what else does the engine expect" - `manuscript.py
# completeness` reports every one of them as FILLED, EMPTY or MISSING.
#
# TWO NAMES THAT ARE EASY TO GET WRONG, and both were wrong on this deck:
#   data/methods_facts.yml   is YAML, not Markdown. There is no methods_facts.md.
#   data/raw_images/         is the folder. There is no source_images/.
#
# AND ONE THAT MOVED: the author list and the affiliations are in
# plan/author_information/, not in the source text. They used to be in
# drafts/source_text/, which is renamed forward on every round - so a paper
# carried one copy of its byline per round, all of them meant to be identical.
# The folder around them is where the signed COI forms go.
RESEARCH_WORK = [
    ("data/methods_facts.yml",
     "Instrument settings, one block per method - the facts the methods "
     "section is built from. It is YAML: there is no methods_facts.md. "
     "scaffold.py prefill writes the skeleton for your field; scaffold.py "
     "harvest reads settings straight off the image metadata."),
    ("plan/outline.md",
     "The # Notes region is YOURS - write what you want said, in any order. "
     "The engine turns it into # Structure and never edits your half. A "
     "complete outline is left alone."),
    ("plan/README.md",
     "The hypothesis and the story. This is the file every blind module is "
     "denied, which is what makes their reports worth reading."),
    ("drafts/rough_draft.md",
     "Optional, and an empty one is a normal project. Write the paper "
     "yourself here and the engine EDITS instead of composing - it is the top "
     "of the content ladder, and it is primary or ignored, never halfway."),
    ("data/raw_images/",
     "Images straight off the instrument - micrographs, gels, spectra. If you "
     "have been calling this source_images/, this is it; the engine does not "
     "look for that name."),
    ("data/raw/",
     "The numbers, as tables. This is what analysis.R reads and what a "
     "figure script points at when the real data lands."),
    ("data/analysis/analysis.R  ->  analysis.md",
     "You fill in two sections of the script - which columns are outcomes, "
     "which split them - and it writes analysis.md. That file is the ONLY "
     "place the drafter may take a number from. The `analysis` skill fills "
     "in those two sections with you."),
    ("plan/captions.md",
     "One caption per float, each with the claim line the figure is supposed "
     "to carry. reviewer-check reads the rendered float against that line."),
    ("drafts/references.bib",
     "Verified through scholar.py, never from memory. citation-check runs in "
     "every preset that writes prose and is never offered as a choice."),
    ("plan/author_information/",
     "authors.md and affiliations.md - author order, affiliations, funding, "
     "CRediT. They ship blank - fill them in on day one, when you already "
     "know the answers. The engine asks if you did not, but a missing figure "
     "is an email and a missing coauthor is a correction notice. "
     "conflict_statements/ beside them takes the signed COI forms."),
]

# The writing engine as the user reads it: six boxes, one round, and it runs
# again. The names in quotes are the plain-English ones; underneath each is
# the module the engine actually runs, because those are what the plan line
# and the reports will say.
WRITING_CHAIN = [
    ("Input", "", [
        "plan/outline.md   plan/README.md",
        "plan/captions.md + claim lines",
        "figures + tables, rendered",
        "data/analysis/analysis.md",
        "data/methods_facts.yml",
        "plan/literature_landscape.md",
        "drafts/references.bib",
        "drafts/rough_draft.md",
    ]),
    ("Drafter", "draft-sections [agent]", [
        "abstract [BLIND]",
        "Writes the body from the outline,",
        "then the title and abstract LAST -",
        "and the abstract agent is denied",
        "the README and the hypothesis.",
    ]),
    ("Organizer", "revise-prose [BLIND]", [
        "assemble",
        "The house voice, denied every",
        "claim source; then pandoc -> the",
        ".docx behind six gates.",
    ]),
    ("Reviewer", "evidence-check", [
        "stats-check [BLIND]",
        "reviewer-check [BLIND]",
        "Every number traces to something",
        "recorded; the built paper is read",
        "the way a reviewer would read it.",
    ]),
    ("Citations Checker", "citation-check [agent]", [
        "Every in-text key resolves, every",
        "entry is cited, no duplicates,",
        "under the cap. Never offered as a",
        "choice - it simply runs.",
    ]),
    ("User Edits", "docx_edits.py -> ingest", [
        "Tracked changes from you and from",
        "every coauthor, attributed per",
        "author, into edits_status.md -",
        "one row per point, nothing lost.",
    ]),
]

# The writing engine, in the order it runs. `mark` drives the legend:
#   ""       an engine call, or a module that reads reports
#   "agent"  spawned as its own sub-agent
#   "blind"  spawned, AND denied the hypothesis - breaking that fails silently
#
# FIFTEEN, not fourteen. `abstract` is a real module - agent-brief lists it,
# it is spawned, and it is one of the FOUR blind ones. It was missing from
# this table, which made the slide understate the thing the slide exists to
# say.
ENGINE_MODULES = [
    ("outline", "agent", "proposes or completes the paragraph plan"),
    ("literature-landscape", "", "what the field says, from six indexes"),
    ("learn-from-edits", "", "your own edits become rules for this draft"),
    ("draft-sections", "agent", "writes the body from the outline"),
    ("abstract", "blind", "title and abstract, denied the README"),
    ("revise-prose", "blind", "the house voice, denied every claim source"),
    ("evidence-check", "", "every number traces to something recorded"),
    ("stats-check", "blind", "denied the hypothesis, on purpose"),
    ("assemble", "", "pandoc -> the .docx, behind six gates"),
    ("citation-check", "agent", "every key resolves; never offered, it runs"),
    ("reviewer-check", "blind", "reads the built paper as a reviewer would"),
    ("final-check", "", "every report, resolved or waived"),
    ("submission-package", "", "refuses while any FLAG remains"),
    ("respond-to-reviewers", "", "one response per point, from the ledger"),
    ("toc-graphic", "", "the graphical abstract, as a slide you draw"),
]

# Everything useful that is not one of the four stages. Titles are the label
# on the box; the command is what it really is.
OTHER_TOOLS = [
    ("AlphaFold, PDB And EMDB", "structure.py alphafold | pdb | emdb | resolve",
     "A predicted or measured structure, fetched and cited properly, ready to "
     "compose into a figure. It is never gap evidence - a structure existing "
     "does not mean the question has been asked."),
    ("The Graphical Abstract", "toc-graphic  /  toc_graphic.py",
     "Sized the journal's own way, as a slide you draw. It runs BEFORE "
     "assemble: the block it writes into source_text/ has to exist before the "
     ".docx is built, or the graphic is simply absent from the submission."),
    ("Drawn Panels", "create-graphic-figure  /  graphic_figure.py add --figure 2",
     "A reaction scheme, an apparatus, a mechanism, a protein figure - drawn "
     "in PowerPoint and composed beside the plotted panels, so one figure can "
     "be half art and half chart and still render as one file."),
    ("Figure Surgery", "refine-figure",
     "For when the plot does not show the claim its caption asserts, or the "
     "axes are wrong, or the mock data does not exercise the hypothesis. Also "
     "lints the float scripts for the traps that render fine and lie."),
    ("Literature", "scholar.py search | verify | check-refs   ·   pubmed.py sections",
     "Six indexes rather than one, because PubMed cannot see half of this "
     "group's field. pubmed.py sections <PMID> pulls the full text of an "
     "open-access paper when the abstract is not enough."),
    ("Flags", "flag-resolver",
     "Walks the outstanding [FLAG: ...] markers one at a time, explains what "
     "each is asking and why it blocks, and records what you decide so the "
     "next round applies it. Runs standalone, on any project, at any time."),
    ("Coauthor Edits", "docx_edits.py extract <file>",
     "Who changed what, per author, out of a tracked-changes .docx - before "
     "any of it becomes a rule. This is the front half of the learning loop "
     "and it is worth running on its own."),
    ("Defects", "manuscript.py log-issue",
     "The engine's own complaints about itself, carried upstream to whoever "
     "maintains the toolkit. Never your prose, never your project name, never "
     "your learned rules. Off until you configure it."),
]

# The two feedback loops. Drawn as arrows going back, because both of them are
# how the toolkit stops repeating itself.
LOOPS = [
    ("Learning Loop", "your edits", "learn.py", "the next draft",
     "A coauthor's tracked change is read by docx_edits.py, classified by "
     "learn.py, and lands in the drafting brief BEFORE the next round writes "
     "a word. Two observations before anything becomes a rule; your registry "
     "never leaves your machine."),
    ("Defect Loop", "the engine catches itself", "log-issue", "system-changes.md",
     "The engine keeps its own defect list. A run that trips over its own "
     "limitation files the defect, in generic terms, with a way to check it "
     "again - and the fix is made there and then. Nothing about your work "
     "leaves the machine it ran on."),
]

# What a sentence in the chat routes to. specs/orchestration.md 2.2 is the
# authority; this is the short version that fits on a slide. The left column
# is QUOTED SPEECH and stays in the casing a person would type it in.
ROUTES = [
    ("“is there a gap here?”", "idea-generation"),
    ("“set up a folder for this paper”", "setup-project-directory"),
    ("“I need a protein figure”", "create-graphic-figure  (+ structure.py)"),
    ("“is anything missing from my folder?”", "scaffold.py check"),
    ("“fake data so I can test a figure”", "idea.py mock"),
    ("“is this citation real?”", "scholar.py verify"),
    ("“draft the introduction”", "writing-engine, scoped"),
    ("“format this for Langmuir”", "writing-engine + format-check"),
    ("“my coauthor sent edits”", "docx_edits.py -> ingest"),
    ("“is this ready to submit?”", "submission-package"),
    ("“this figure is ugly / wrong”", "refine-figure"),
    ("“what does AlphaFold have?”", "structure.py  (never gap evidence)"),
    ("“the engine did something wrong”", "manuscript.py log-issue"),
]

MODULAR = [
    ("Scaffold A Project", "scaffold.py scaffold <path>"),
    ("Check One For Missing Pieces", "scaffold.py check <path>"),
    ("Mock Data / Mock Floats", "idea.py mock  /  scaffold.py mock-floats"),
    ("Search Or Verify Literature", "scholar.py search / verify / check-refs"),
    ("Full Text Of An Open-Access Paper", "pubmed.py sections <PMID>"),
    ("Art Panel For A Figure", "graphic_figure.py add --figure 2"),
    ("Who Changed What In A .docx", "docx_edits.py extract <file>"),
    ("Build The Manuscript", "manuscript.py assemble <path> --journal X"),
    ("What Is Still Missing", "manuscript.py completeness <path>"),
    ("Number Density / Callouts / Length", "prose.py density / crossrefs / length"),
    ("What It Has Learned From You", "learn.py brief"),
    ("Lint The Float Scripts", "scaffold.py float lint <path>"),
    ("Prefill Instrument Settings", "scaffold.py prefill <path>"),
    ("Read The Settings Off The Images", "scaffold.py harvest <path>"),
    ("Merge An Idea Into A Project", "idea.py merge <path> --bundle b.json"),
    ("A Predicted Or Measured Structure",
     "structure.py alphafold / pdb / emdb"),
    ("The AI-Use Declaration", "manuscript.py ai-disclosure <path>"),
]


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------

def _rgb(hexstr: str) -> "RGBColor":
    return RGBColor.from_string(hexstr)


def _text(frame, lines, size=10, color=INK, bold_first=False, align=PP_ALIGN.LEFT,
          space_after=2):
    """Fill a text frame with one paragraph per line.

    `lines` items are either a string or (text, size, bold, color).
    """
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.05)
    frame.margin_bottom = Inches(0.05)
    first = True
    for item in lines:
        if isinstance(item, tuple):
            txt, sz, bold, col = item
        else:
            txt, sz, bold, col = item, size, (bold_first and first), color
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        run = para.add_run()
        run.text = txt
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = _rgb(col)
        run.font.name = "Segoe UI"
        first = False


def _box(slide, x, y, w, h, fill, line, lines, size=10, radius=0.06,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.0, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(fill)
    sh.line.color.rgb = _rgb(line)
    sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        # The default corner is a caricature at this size.
        try:
            sh.adjustments[0] = radius
        except (IndexError, KeyError):
            pass
    tf = sh.text_frame
    tf.vertical_anchor = anchor
    _text(tf, lines, size=size, align=align)
    return sh


def _label(slide, x, y, w, h, lines, size=10, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _text(sh.text_frame, lines, size=size, align=align)
    return sh


def _arrow(slide, x, y, w, h, color=MUTED, down=False, left=False):
    shape = MSO_SHAPE.DOWN_ARROW if down else (
        MSO_SHAPE.LEFT_ARROW if left else MSO_SHAPE.RIGHT_ARROW)
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _rule(slide, x, y, w, color=RULE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Emu(9525))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _heading(slide, title, sub=""):
    _label(slide, 0.55, 0.30, 12.2, 0.5,
           [(title, 26, True, INK)], size=26)
    if sub:
        _label(slide, 0.58, 0.80, 12.2, 0.35, [(sub, 12, False, MUTED)], size=12)
    _rule(slide, 0.55, 1.18, 12.2)


def _footer(slide, text):
    _label(slide, 0.55, 7.02, 12.2, 0.3, [(text, 9, False, MUTED)], size=9)


# ---------------------------------------------------------------------------
# Animation: the <p:timing> block
# ---------------------------------------------------------------------------
# One click per step. Each step is a list of shape ids that appear together:
# the first is a clickEffect, the rest are withEffect, which is exactly what
# PowerPoint itself writes for "appear, with previous".
#
# The nesting below (seq -> par -> par -> par -> effect) is not decoration.
# CT_TLTimeNodeParallel is what each level is, and PowerPoint rejects the file
# rather than flattening it, so it is reproduced as PowerPoint writes it.

_EFFECT = (
    '<p:par><p:cTn id="{cid}" presetID="1" presetClass="entr" presetSubtype="0"'
    ' fill="hold" grpId="0" nodeType="{node}">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:childTnLst>'
    '<p:set>'
    '<p:cBhvr>'
    '<p:cTn id="{cid2}" dur="1" fill="hold">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '</p:cTn>'
    '<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
    '</p:cBhvr>'
    '<p:to><p:strVal val="visible"/></p:to>'
    '</p:set>'
    '</p:childTnLst>'
    '</p:cTn></p:par>'
)

_STEP = (
    '<p:par><p:cTn id="{a}" fill="hold">'
    '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
    '<p:childTnLst>'
    '<p:par><p:cTn id="{b}" fill="hold">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
    '<p:childTnLst>{effects}</p:childTnLst>'
    '</p:cTn></p:par>'
    '</p:childTnLst>'
    '</p:cTn></p:par>'
)

_TIMING = (
    '<p:timing xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:tnLst>'
    '<p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
    '<p:childTnLst>'
    '<p:seq concurrent="1" nextAc="seek">'
    '<p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
    '<p:childTnLst>{steps}</p:childTnLst>'
    '</p:cTn>'
    '<p:prevCondLst><p:cond evt="onPrev" delay="0">'
    '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
    '<p:nextCondLst><p:cond evt="onNext" delay="0">'
    '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
    '</p:seq>'
    '</p:childTnLst>'
    '</p:cTn></p:par>'
    '</p:tnLst>'
    '<p:bldLst>{builds}</p:bldLst>'
    '</p:timing>'
)


def _animate(slide, steps) -> int:
    """Give `slide` one click per step. Returns the number of effects written.

    `steps` holds shape objects, grouped: everything in one group appears on
    one click.
    """
    groups = [[sh for sh in group if sh is not None] for group in steps]
    groups = [g for g in groups if g]
    if not groups:
        return 0

    cid = 3          # 1 is tmRoot and 2 is mainSeq
    step_xml = []
    builds = []
    effects_written = 0
    for group in groups:
        a, b = cid, cid + 1
        cid += 2
        eff = []
        for i, sh in enumerate(group):
            eff.append(_EFFECT.format(
                cid=cid, cid2=cid + 1, spid=sh.shape_id,
                node="clickEffect" if i == 0 else "withEffect"))
            cid += 2
            builds.append('<p:bldP spid="%d" grpId="0" animBg="1"/>' % sh.shape_id)
            effects_written += 1
        step_xml.append(_STEP.format(a=a, b=b, effects="".join(eff)))

    xml = _TIMING.format(steps="".join(step_xml), builds="".join(builds))
    # Parsed rather than concatenated, so a malformed block fails here - in the
    # build - instead of in PowerPoint on somebody else's machine.
    node = etree.fromstring(xml)
    # <p:timing> is the last child of <p:sld>; anything already there is a
    # rebuild and is replaced rather than appended to.
    sld = slide._element
    for old in sld.findall(qn("p:timing")):
        sld.remove(old)
    sld.append(node)
    return effects_written


# ---------------------------------------------------------------------------
# The slides
# ---------------------------------------------------------------------------

def _blank(prs):
    # Layout 6 of the default template is the blank one. Everything here is
    # drawn, so a layout with placeholders would put empty prompt boxes on
    # every slide.
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide_title(prs, animate):
    s = _blank(prs)
    _label(s, 0.9, 1.9, 11.5, 1.0, [("Paper Engine", 46, True, INK)], size=46)
    _label(s, 0.95, 2.95, 11.0, 0.6,
           [("How The Writing Aids Fit Together", 18, False, MUTED)], size=18)
    _rule(s, 0.95, 3.55, 6.0)
    steps = []
    x = 0.95
    for st in STAGES:
        b = _box(s, x, 3.95, 2.75, 1.15, STAGE_FILL[st["key"]],
                 STAGE_LINE[st["key"]],
                 [(st["title"], 13, True, INK), (st["sub"], 9, False, MUTED)],
                 align=PP_ALIGN.CENTER)
        steps.append([b])
        x += 3.0
    _label(s, 0.95, 5.45, 11.5, 0.9,
           [("Four stages, and the order is a default rather than a rule - you "
             "can arrive at any of them.", 11, False, MUTED),
            ("Nothing here writes anything without showing you first.",
             11, False, MUTED)], size=11)
    _footer(s, "specs/orchestration.md  |  help/instructions.md")
    if animate:
        _animate(s, steps)
    return s


def slide_overview(prs, animate):
    s = _blank(prs)
    _heading(s, "The Whole Flow",
             "One block per stage; the arrows are hand-offs. Two more arrows "
             "go backwards, and they are drawn on the writing-engine slide.")
    steps = []
    y = 1.45
    for st in STAGES:
        k = st["key"]
        box = _box(s, 0.55, y, 3.5, 1.28, STAGE_FILL[k], STAGE_LINE[k],
                   [("Stage %s  -  %s" % (k, st["title"]), 13, True, INK),
                    (st["sub"], 9, False, MUTED)])
        agents = _box(s, 4.35, y, 4.35, 1.28, PAPER, RULE,
                      [(nm, 10, True, INK) for nm, _ in st["agents"]], size=10)
        engines = _box(s, 9.0, y, 3.75, 1.28, ENGINE_FILL, ENGINE_LINE,
                       [("Engines", 8, True, MUTED),
                        ("  ".join(st["engines"]), 9, False, INK)], size=9)
        arrow = None
        if k != "D":
            arrow = _arrow(s, 2.18, y + 1.29, 0.20, 0.15, down=True,
                           color=STAGE_LINE[k])
        steps.append([box, agents, engines, arrow])
        y += 1.42

    loop = _box(s, 4.35, 7.02, 8.4, 0.42, PAPER, LOOP_LINE,
                [("Two arrows go back: your edits become drafting rules "
                  "(learn.py), and the engine's own defects become items in "
                  "system-changes.md (log-issue).", 9, False, INK)], size=9)
    steps.append([loop])
    if animate:
        _animate(s, steps)
    return s


def slide_idea(prs, animate):
    """Stage A as five boxes in the order the engine really runs them.

    The order is the point of this slide. People list it as articles, talk,
    mock data, figures, folder - and the folder has to come third, because
    `idea.py mock` writes into `data/mock_data/` and the draft figures are
    rendered FROM that mock data by `scaffold.py mock-floats`. Drawing it in
    the order people say it would put two steps in front of the thing they
    both need.
    """
    s = _blank(prs)
    _heading(s, "Idea Generation",
             "Five steps, in the order the engine actually runs them - ending "
             "in a folder that already renders")
    steps = []

    y = 1.42
    for i, (title, body) in enumerate(IDEA_STEPS):
        b = _box(s, 0.55, y, 6.15, 0.95, STAGE_FILL["A"], STAGE_LINE["A"], [
            ("%d.  %s" % (i + 1, title), 12, True, INK),
            (body, 8, False, INK),
        ], size=8)
        arrow = None
        if i < len(IDEA_STEPS) - 1:
            arrow = _arrow(s, 3.50, y + 0.97, 0.24, 0.16, down=True,
                           color=STAGE_LINE["A"])
        steps.append([b, arrow])
        y += 1.13

    # The image slot. Left empty on purpose - this is where a screenshot of a
    # real project tree goes, and a drawn approximation of one would be a
    # picture that goes stale without anybody noticing.
    slot = _box(s, 7.00, 1.42, 5.78, 4.62, PAPER, MUTED, [
        ("", 9, False, INK),
        ("[  Paste A Screenshot Of An Example  ]", 14, True, MUTED),
        ("[  Project Directory Here  ]", 14, True, MUTED),
        ("", 9, False, INK),
        ("Right-click this box  ->  Format Shape  ->  Fill  ->  "
         "Picture, or just paste over it.", 9, False, MUTED),
    ], size=9, align=PP_ALIGN.CENTER)
    steps.append([slot])

    tree = _box(s, 7.00, 6.15, 5.78, 0.92, ENGINE_FILL, ENGINE_LINE, [
        ("What The Scaffold Writes", 10, True, INK),
        ("plan/  (outline, README, captions, author_information/, "
         "figures/FigNN/, tables/, theme/)  "
         "·  data/  (raw/, raw_images/, analysis/, mock_data/, "
         "methods_facts.yml, data_contract.md)  ·  drafts/  (rough_draft.md, "
         "source_text/, edits/, references.bib)  ·  obsolete/  ·  project.yml",
         8, False, INK),
    ], size=8)
    steps.append([tree])

    _footer(s, "Every citation goes through scholar.py verify before it "
               "reaches a file. No citation ever comes from memory.")
    if animate:
        _animate(s, steps)
    return s


def slide_research_work(prs, animate):
    """Stage B, as the list of files the writing engine will actually read.

    This is WEIGHTABLE_SOURCES in manuscript.py plus the two files
    submission-package needs, which is the only honest answer to "what else
    does the engine expect" - anything else is a guess, and `manuscript.py
    completeness` reports exactly these.
    """
    s = _blank(prs)
    _heading(s, "Research Work - What You Fill In",
             "The aids hold the shape of the folder and stay out of it. These "
             "are the files the writing engine reads.")
    steps = []

    for i, (path, body) in enumerate(RESEARCH_WORK):
        x = 0.55 + (i % 2) * 6.42
        y = 1.40 + (i // 2) * 0.99
        b = _box(s, x, y, 6.20, 0.92, STAGE_FILL["B"], STAGE_LINE["B"], [
            (path, 11, True, INK),
            (body, 8, False, INK),
        ], size=8, anchor=MSO_ANCHOR.TOP)
        steps.append([b])

    mock = _box(s, 0.55, 6.34, 7.35, 0.82, BLIND_FILL, "A9601F", [
        ("Mock Data Is A Head Start, Not A Requirement", 10, True, INK),
        ("data/mock_data/*_mock.csv is there from day one so every figure "
         "script runs before any data does. Adjust the generator as the design "
         "settles, re-render, and point the script at data/raw/ when the real "
         "numbers land - which is also how the MOCK watermark goes away.",
         8, False, INK),
    ], size=8, anchor=MSO_ANCHOR.TOP)
    steps.append([mock])

    ask = _box(s, 8.15, 6.34, 4.60, 0.82, ENGINE_FILL, ENGINE_LINE, [
        ("Do Not Guess - Ask", 10, True, INK),
        ("manuscript.py completeness \"<project>\" reports every file above "
         "as FILLED, EMPTY or MISSING. EMPTY and MISSING are different "
         "problems: an EMPTY file exists and needs writing in, a MISSING one "
         "has to be created first.", 8, False, INK),
    ], size=8, anchor=MSO_ANCHOR.TOP)
    steps.append([ask])

    if animate:
        _animate(s, steps)
    return s


def slide_writing_engine(prs, animate):
    """Stage C as one round, drawn left to right, with both arrows back.

    Each box carries the plain-English name AND the module the engine really
    runs, because the plan line, the reports and every error message will use
    the second one.
    """
    s = _blank(prs)
    _heading(s, "The Writing Engine",
             "One round, left to right. Then it runs again - and every round "
             "it knows more about how you write.")
    steps = []

    bw, gap = 1.84, 0.22
    x0 = 0.60
    xs = [x0 + i * (bw + gap) for i in range(len(WRITING_CHAIN))]

    # Target journal requirements come in at the Organizer and nowhere else -
    # `assemble` is the module that reads requirements.yml.
    org_x = xs[2]
    req = _box(s, org_x - 0.38, 1.32, 2.60, 0.60, STAGE_FILL["D"],
               STAGE_LINE["D"], [
                   ("Target Journal Requirements", 9, True, INK),
                   ("journal_requirements/requirements.yml", 7, False, MUTED),
               ], size=7, align=PP_ALIGN.CENTER)
    req_arrow = _arrow(s, org_x + bw / 2 - 0.11, 1.96, 0.22, 0.28, down=True,
                       color=STAGE_LINE["D"])
    steps.append([req, req_arrow])

    for i, (title, module, body) in enumerate(WRITING_CHAIN):
        first = i == 0
        fill = ENGINE_FILL if first else STAGE_FILL["C"]
        line = ENGINE_LINE if first else STAGE_LINE["C"]
        lines = [(title, 11, True, INK)]
        if module:
            lines.append((module, 8, True, STAGE_LINE["C"]))
        lines += [(t, 7, False, INK) for t in body]
        b = _box(s, xs[i], 2.28, bw, 1.72, fill, line, lines, size=7,
                 anchor=MSO_ANCHOR.TOP)
        arrow = None
        if i < len(WRITING_CHAIN) - 1:
            arrow = _arrow(s, xs[i] + bw + 0.02, 3.03, 0.18, 0.22,
                           color=STAGE_LINE["C"])
        steps.append([b, arrow])

    # Out of the round and into the submission.
    down = _arrow(s, xs[-1] + bw / 2 - 0.11, 4.02, 0.22, 0.34, down=True,
                  color=STAGE_LINE["D"])
    subm = _box(s, 9.35, 4.40, 3.40, 0.78, STAGE_FILL["D"], STAGE_LINE["D"], [
        ("Submission", 12, True, INK),
        ("final-check  ->  submission-package. Refuses while any [FLAG: ...] "
         "remains anywhere in the paper.", 8, False, INK),
    ], size=8)
    steps.append([down, subm])

    why = _box(s, 0.60, 4.40, 8.55, 0.78, PAPER, LOOP_LINE, [
        ("BLIND Is The Load-Bearing Word", 10, True, LOOP_LINE),
        ("Four modules are given an explicit file list and denied the "
         "hypothesis, the outline and plan/README.md. An agent that knows what "
         "you hoped to find reads an ambiguous result as favourable - and "
         "nothing visibly fails when it does. The report still looks fine.",
         8, False, INK),
    ], size=8)
    steps.append([why])

    # Arrow one: the round runs again.
    back1 = _arrow(s, 0.75, 5.36, 10.90, 0.26, color=LOOP_LINE, left=True)
    lab1 = _label(s, 0.60, 5.66, 12.20, 0.40, [
        ("Iterative - the whole round runs again.", 10, True, LOOP_LINE),
        ("Your edits come back in through ingest, the frozen sections are left "
         "alone, and every checker re-runs on what actually changed. Agent "
         "calls per preset: draft 7  |  sections 7  |  revision 8  |  "
         "coauthor 10  |  submission 13 - one fewer whenever the abstract came "
         "from your own rough draft. The plan is always stated and always "
         "waits for you.",
         8, False, INK),
    ], size=8)
    steps.append([back1, lab1])

    # Arrow two: and it gets better at drafting while it does.
    back2 = _arrow(s, 2.66, 6.42, 8.99, 0.26, color=LOOP_LINE, left=True)
    lab2 = _label(s, 0.60, 6.72, 12.20, 0.45, [
        ("And it learns - the second arrow is why the drafter improves.",
         10, True, LOOP_LINE),
        ("docx_edits.py reads your tracked changes, learn.py classifies them "
         "(two observations before anything becomes a rule), and "
         "learn-from-edits puts them in the drafting brief BEFORE the next "
         "round writes a word. Your registry never leaves your machine.",
         8, False, INK),
    ], size=8)
    steps.append([back2, lab2])

    if animate:
        _animate(s, steps)
    return s


def slide_engine_modules(prs, animate):
    s = _blank(prs)
    _heading(s, "Inside The Writing Engine",
             "All fifteen modules, in the order they run. Run all of it, one "
             "section, or one module.")
    steps = []
    for i, (name, mark, body) in enumerate(ENGINE_MODULES):
        col = i % 4
        row = i // 4
        x = 0.55 + col * 3.09
        y = 1.42 + row * 1.16
        fill = BLIND_FILL if mark == "blind" else (
            PAPER if mark == "agent" else ENGINE_FILL)
        line = LOOP_LINE if mark == "blind" else (
            STAGE_LINE["C"] if mark == "agent" else ENGINE_LINE)
        tag = {"agent": "Own Context", "blind": "BLIND", "": "Engine Call"}[mark]
        box = _box(s, x, y, 2.90, 1.08, fill, line, [
            ("%d.  %s" % (i + 1, name), 11, True, INK),
            (tag, 8, True, MUTED),
            (body, 8, False, INK),
        ], size=8)
        arrow = None
        if col < 3 and i + 1 < len(ENGINE_MODULES):
            arrow = _arrow(s, x + 2.92, y + 0.44, 0.14, 0.20,
                           color=STAGE_LINE["C"])
        steps.append([box, arrow])

    legend = _box(s, 0.55, 6.14, 12.20, 1.06, PAPER, RULE, [
        ("Seven Modules Are Spawned With Their Own Context. Four Of Those Are "
         "Blind.", 11, True, LOOP_LINE),
        ("A blind module is handed an explicit list of files and denied the "
         "hypothesis, the outline and plan/README.md. stats-check is the one "
         "that matters most: an agent that knows what you hoped to find reads "
         "an ambiguous result as favourable - and nothing visibly fails when "
         "that happens. The report still looks fine. It is just worthless. "
         "The plan line names every isolated and blind module before the round "
         "runs, because whether a module ran blind cannot be seen in its "
         "report afterwards.", 8, False, INK),
    ], size=8)
    steps.append([legend])
    if animate:
        _animate(s, steps)
    return s


def slide_other_tools(prs, animate):
    s = _blank(prs)
    _heading(s, "Other Tools That Help",
             "Callable from a skill or straight from a terminal, in the middle "
             "of anything else")
    steps = []
    for i, (title, cmd, body) in enumerate(OTHER_TOOLS):
        x = 0.55 + (i % 2) * 6.42
        y = 1.48 + (i // 2) * 1.32
        b = _box(s, x, y, 6.20, 1.22, PAPER, STAGE_LINE["A"], [
            (title, 12, True, INK),
            (cmd, 8, True, STAGE_LINE["A"]),
            (body, 8, False, INK),
        ], size=8)
        steps.append([b])
    _footer(s, "Each one takes --json, refuses rather than guessing, and needs "
               "nothing from the conversation.")
    if animate:
        _animate(s, steps)
    return s


def slide_routing(prs, animate):
    s = _blank(prs)
    _heading(s, "You Do Not Have To Know The Names",
             "Say what you want; the right module gets offered - offered, "
             "never run")
    steps = []
    for i, (said, goes) in enumerate(ROUTES):
        x = 0.55 + (i % 2) * 6.25
        y = 1.42 + (i // 2) * 0.72
        b = _box(s, x, y, 6.05, 0.66, PAPER, RULE, [
            (said, 10, True, INK),
            ("->  " + goes, 9, False, STAGE_LINE["A"]),
        ], size=9)
        steps.append([b])

    rule = _box(s, 0.55, 6.52, 12.2, 0.68, STAGE_FILL["C"], STAGE_LINE["C"], [
        ("Two Rules Keep This From Being Annoying", 10, True, INK),
        ("A cue names the module and asks - it never fires a skill off a "
         "passing mention. And one skill owns the turn: if you ask for two "
         "things, it says so and starts with the one that has to happen first.",
         9, False, INK),
    ], size=9)
    steps.append([rule])
    if animate:
        _animate(s, steps)
    return s


def slide_modular(prs, animate):
    s = _blank(prs)
    _heading(s, "Everything Is Callable On Its Own",
             "The skills are convenience; these run from a terminal with no "
             "agent at all")
    steps = []
    for i, (ask, cmd) in enumerate(MODULAR):
        y = 1.40 + i * 0.31
        b = _box(s, 0.55, y, 4.60, 0.29, PAPER, RULE, [(ask, 9, False, INK)],
                 size=9)
        c = _box(s, 5.25, y, 7.50, 0.29, ENGINE_FILL, ENGINE_LINE,
                 [("python tools/" + cmd, 9, False, INK)], size=9)
        steps.append([b, c])

    note = _box(s, 0.55, 6.80, 12.2, 0.45, STAGE_FILL["B"], STAGE_LINE["B"], [
        ("Each one takes --json, refuses rather than guessing, and needs "
         "nothing from the conversation. That is what makes it safe to call one "
         "in the middle of something else.", 9, False, INK),
    ], size=9)
    steps.append([note])
    if animate:
        _animate(s, steps)
    return s


SLIDES = [slide_title, slide_overview,
          slide_idea, slide_research_work,
          slide_writing_engine, slide_engine_modules,
          slide_other_tools,
          slide_routing, slide_modular]


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------

def build(out: str, animate: bool = True) -> dict:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for fn in SLIDES:
        fn(prs, animate)
    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    prs.save(out)
    res = {"path": os.path.abspath(out), "slides": len(SLIDES),
           "animated": animate}
    res.update(check(out))
    return res


def check(path: str) -> dict:
    """Re-open the written file and measure it.

    Every XML part is re-parsed, because the one failure this generator can
    produce - a malformed timing block - is invisible until PowerPoint refuses
    the file, and by then it is on somebody else's machine.
    """
    out: dict = {"checked": os.path.abspath(path), "problems": []}
    if not os.path.isfile(path):
        out["problems"].append("no such file")
        return out

    with zipfile.ZipFile(path) as z:
        parts = z.namelist()
        for name in parts:
            if not name.endswith((".xml", ".rels")):
                continue
            try:
                etree.fromstring(z.read(name))
            except etree.XMLSyntaxError as exc:
                out["problems"].append("%s: %s" % (name, exc))
        slides = sorted(p for p in parts
                        if p.startswith("ppt/slides/slide") and p.endswith(".xml"))
        out["slide_parts"] = len(slides)
        effects = 0
        timed = 0
        for name in slides:
            root = etree.fromstring(z.read(name))
            tim = root.findall(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}timing")
            if tim:
                timed += 1
            effects += len(root.findall(
                ".//{http://schemas.openxmlformats.org/presentationml/2006/main}set"))
        out["slides_with_timing"] = timed
        out["effects"] = effects

    prs = Presentation(path)
    out["slides"] = len(prs.slides)
    out["shapes"] = sum(len(sl.shapes) for sl in prs.slides)

    # EVERY SHAPE HAS TO BE ON THE SLIDE IT IS ON.
    # Nothing here measured where a shape ENDS, and two slides had been
    # running off the bottom edge for as long as their lists had been that
    # long - the routing slide by half an inch, the modular slide by two
    # inches, with its last five commands simply not in the deck. Both were
    # written as `y = top + i * step` with no assertion, both opened fine, and
    # both looked authoritative. That is the same failure mode as a stale
    # figure, so it gets the same treatment: measured, not trusted.
    #
    # A quarter-inch of slack, because a rounded corner's bounding box and a
    # deliberate full-bleed rule are not the bug this is looking for.
    slack = Inches(0.25)
    # A presentation can legitimately declare no slide size, in which case
    # there is no edge to be off and the loop below simply does not run.
    sw = prs.slide_width or 0
    sh_ = prs.slide_height or 0
    bounded = list(enumerate(prs.slides, 1)) if (sw and sh_) else []
    for i, sl in bounded:
        for shp in sl.shapes:
            left, top = shp.left, shp.top
            right = left + (shp.width or 0)
            bottom = top + (shp.height or 0)
            if (left < -slack or top < -slack
                    or right > sw + slack or bottom > sh_ + slack):
                out["problems"].append(
                    "slide %d: %r runs off the slide "
                    "(left %.2f top %.2f right %.2f bottom %.2f, "
                    "slide is %.2f x %.2f inches)"
                    % (i, (shp.name or "shape")[:40],
                       left / 914400, top / 914400,
                       right / 914400, bottom / 914400,
                       sw / 914400, sh_ / 914400))
    # Every animated shape id must exist on its own slide, or PowerPoint drops
    # the effect silently - the animation is gone and the deck still opens,
    # which is the worst of the failure modes available here.
    ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for i, sl in enumerate(prs.slides, 1):
        ids = {sh.shape_id for sh in sl.shapes}
        for tgt in sl._element.findall(".//%sspTgt" % ns):
            spid = int(tgt.get("spid") or 0)
            if spid not in ids:
                out["problems"].append(
                    "slide %d animates shape %d, which is not on it"
                    % (i, spid))
    out["ok"] = not out["problems"]
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=(__doc__ or "").splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="write the deck")
    b.add_argument("--out", default=DEFAULT_OUT)
    b.add_argument("--no-animation", action="store_true",
                   help="write the same deck with no <p:timing> block")
    b.add_argument("--json", action="store_true")

    c = sub.add_parser("check", help="re-open a built deck and measure it")
    c.add_argument("--out", default=DEFAULT_OUT)
    c.add_argument("--json", action="store_true")

    g = sub.add_parser("graph", help="the stage graph, as data")
    g.add_argument("--json", action="store_true")

    args = ap.parse_args(argv)

    if args.cmd == "graph":
        data = {"stages": STAGES, "engine_modules": ENGINE_MODULES,
                "loops": [dict(zip(("title", "from", "via", "to", "why"), L))
                          for L in LOOPS],
                "routes": ROUTES, "modular": MODULAR}
        print(json.dumps(data, indent=2) if args.json else
              "\n".join("%s  %s" % (s["key"], s["title"]) for s in STAGES))
        return 0

    if args.cmd == "build":
        res = build(args.out, animate=not args.no_animation)
    else:
        res = check(args.out)

    if args.json:
        print(json.dumps(res, indent=2))
    else:
        if res.get("error"):
            print("REFUSED - " + res["error"])
            return 2
        print("  %s" % res.get("path", res.get("checked")))
        print("  slides: %s   shapes: %s   animated slides: %s   effects: %s"
              % (res.get("slides"), res.get("shapes"),
                 res.get("slides_with_timing"), res.get("effects")))
        for p in res.get("problems", []):
            print("  PROBLEM  " + p)
        print("  %s" % ("ok" if res.get("ok") else "PROBLEMS FOUND"))
    return 0 if res.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main())
